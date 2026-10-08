from copy import deepcopy
from dataclasses import dataclass, field

from docx.document import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml.shared import OxmlElement
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from ..compilers.docx_utils import paragraph_text
from ..i18n import t
from .content_interface import IDocumentContent, new_content_id
from .model import IContent
from .rich_text import cell_value, parse_cell_value, set_cell_value
from .styles import caption_style, find_style

CAPTION_POSITIONS = ("above", "below")
TABLE_ALIGNMENTS = {"left": WD_TABLE_ALIGNMENT.LEFT, "center": WD_TABLE_ALIGNMENT.CENTER, "right": WD_TABLE_ALIGNMENT.RIGHT}
ALIGNMENT_NAMES = {value: name for name, value in TABLE_ALIGNMENTS.items()}


@dataclass
class WordTable(IContent, IDocumentContent):
    """
    Word table with its content and caption.
    rows: cells by row (merged cells appear once): plain text, or {"html": ...} for rich text with links (see
          rich_text);
    caption: caption text ("" if missing), above or below the table according to caption_position;
    header: the first row is a header (bold and repeated on every page);
    style: id of the table style; alignment: "left", "center", "right" or None (the style's one).
    """
    rows: list[list[str | dict]]
    caption: str = ""
    caption_position: str = "below"
    header: bool = False
    style: str | None = None
    alignment: str | None = None
    id: int | None = None
    anchor: int | None = None
    content_id: str = field(default_factory=new_content_id)

    @classmethod
    def from_docx(cls, table: Table, caption: Paragraph | None = None, caption_position: str = "below") -> "WordTable":
        """Reads content, style, header and caption of an existing table."""
        rows = [[cell_value(_Cell(cell, table)) for cell in row.tc_lst] for row in table._tbl.tr_lst]
        first_row = table._tbl.tr_lst[0] if table._tbl.tr_lst else None
        header = first_row is not None and first_row.find(f"{qn('w:trPr')}/{qn('w:tblHeader')}") is not None
        style = table._tbl.tblPr.find(qn("w:tblStyle"))
        return cls(
            rows=rows,
            caption=paragraph_text(caption).strip() if caption is not None else "",
            caption_position=caption_position,
            header=header,
            style=style.get(qn("w:val")) if style is not None else None,
            alignment=ALIGNMENT_NAMES.get(table.alignment),
        )

    @classmethod
    def from_dict(cls, value: dict) -> "WordTable":
        """Creates the table from the request JSON; raises ValueError if a field is not valid."""
        rows = value.get("rows")
        if not isinstance(rows, list) or not rows or not all(isinstance(row, list) and row for row in rows):
            raise ValueError(t("errors.tableRows"))
        rows = [[parse_cell_value(cell) for cell in row] for row in rows]
        caption = value.get("caption") or ""
        caption_position = value.get("caption_position") or "below"
        alignment = value.get("alignment") or None
        if not isinstance(caption, str):
            raise ValueError(t("errors.captionNotString"))
        if caption_position not in CAPTION_POSITIONS:
            raise ValueError(t("errors.captionPosition"))
        if alignment is not None and alignment not in TABLE_ALIGNMENTS:
            raise ValueError(t("errors.unsupportedTableAlignment", alignment=alignment))
        return cls(
            rows=rows,
            caption=caption.strip(),
            caption_position=caption_position,
            header=bool(value.get("header")),
            style=value.get("style") or None,
            alignment=alignment,
        )

    def to_dict(self) -> dict:
        """Type "table" with rows, caption and style."""
        return {
            "type": "table",
            "id": self.id,
            "anchor": self.anchor,
            "layout": self.layout,
            "rows": self.rows,
            "caption": self.caption,
            "caption_position": self.caption_position,
            "header": self.header,
            "style": self.style,
            "alignment": self.alignment,
        }

    def serialize(self) -> dict:
        """Type "table" as a content (e.g. of a ContentBlock): id is the content_id, not the position in the document."""
        return {
            "type": "table",
            "id": self.content_id,
            "rows": self.rows,
            "caption": self.caption,
            "caption_position": self.caption_position,
            "header": self.header,
            "style": self.style,
            "alignment": self.alignment,
        }

    def build(self, document: Document) -> list:
        """
        Creates table and caption at the end of the document and returns their elements in insertion order
        (DocumentBuilder moves them to the requested position).
        """
        style = find_style(document, self.style, WD_STYLE_TYPE.TABLE)
        columns = max(len(row) for row in self.rows)
        # add_table spreads the columns over the usable page width: the table does not overflow.
        table = document.add_table(rows=len(self.rows), cols=columns)
        if style is not None:
            table.style = style
        if self.alignment:
            table.alignment = TABLE_ALIGNMENTS[self.alignment]
        for row, values in zip(table.rows, self.rows):
            for cell, value in zip(row.cells, values):
                if isinstance(value, str):
                    cell.text = value
                else:
                    set_cell_value(cell, value)
        if self.header:
            _make_bold(table.rows[0]._tr)
            table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))

        elements = [table._tbl]
        if self.caption:
            caption = document.add_paragraph(self.caption, style=caption_style(document))
            if self.caption_position == "above":
                # A caption above stays on the same page as the table.
                caption.paragraph_format.keep_with_next = True
                elements.insert(0, caption._p)
            else:
                elements.append(caption._p)
        return elements

    def update(self, document: Document, table: Table, caption: Paragraph | None) -> bool:
        """
        Updates an existing table (and its caption) in place, keeping borders, shading, widths and cell
        formatting. New rows copy the last row. Returns False, without changing anything, if the number of
        columns changes: in that case the table must be recreated with build.
        """
        table_element = table._tbl
        columns = max(len(row.tc_lst) for row in table_element.tr_lst)
        if max(len(row) for row in self.rows) != columns:
            return False

        rows = table_element.tr_lst
        was_header = rows[0].find(f"{qn('w:trPr')}/{qn('w:tblHeader')}") is not None
        # If there is only the header, the new rows copy it without shading and bold.
        copies_header = len(rows) == 1 and (was_header or self.header)
        while len(rows) < len(self.rows):
            new_row = deepcopy(rows[-1])
            tags = [qn("w:tblHeader")] + ([qn("w:shd"), qn("w:b"), qn("w:bCs")] if copies_header else [])
            for element in list(new_row.iter(*tags)):
                element.getparent().remove(element)
            rows[-1].addnext(new_row)
            rows = table_element.tr_lst
        for extra_row in rows[len(self.rows):]:
            table_element.remove(extra_row)
        # Merged cells appear once per row, as in from_docx.
        for row, values in zip(table_element.tr_lst, self.rows):
            for cell_element, value in zip(row.tc_lst, values):
                set_cell_value(_Cell(cell_element, table), value)

        first_row = table_element.tr_lst[0]
        if self.header and not was_header:
            first_row.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            _make_bold(first_row)
        elif not self.header:
            for header_mark in first_row.iter(qn("w:tblHeader")):
                header_mark.getparent().remove(header_mark)
        style = find_style(document, self.style, WD_STYLE_TYPE.TABLE)
        if style is not None:
            table.style = style
        table.alignment = TABLE_ALIGNMENTS.get(self.alignment)

        self._update_caption(document, table_element, caption)
        return True

    def _update_caption(self, document: Document, table_element, caption: Paragraph | None) -> None:
        """Creates, updates, moves or deletes the table caption; if the text does not change it stays as it is (fields included)."""
        if not self.caption:
            if caption is not None:
                caption._p.getparent().remove(caption._p)
            return
        if caption is None:
            caption = document.add_paragraph(style=caption_style(document))
        if paragraph_text(caption).strip() != self.caption:
            for child in list(caption._p):
                if child.tag != qn("w:pPr"):
                    caption._p.remove(child)
            caption.add_run(self.caption)
        if self.caption_position == "above":
            caption.paragraph_format.keep_with_next = True
            table_element.addprevious(caption._p)
        else:
            caption.paragraph_format.keep_with_next = None
            table_element.addnext(caption._p)


def _make_bold(row_element) -> None:
    """Bold for all the runs of the row, also those inside hyperlinks (a header row)."""
    for run_element in row_element.iter(qn("w:r")):
        # Through python-docx, which puts w:b in the position required by the schema.
        Run(run_element, None).font.bold = True
