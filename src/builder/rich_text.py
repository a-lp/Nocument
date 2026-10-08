"""
Rich text of the visual editor (see html_converter) written into Word paragraphs, with real hyperlinks.

- append_html_runs writes the runs of a block (bold, italic, underline, strikethrough, line breaks and links) at the
  end of a paragraph; html_run_elements creates the same elements without placing them (e.g. in place of a keyword).
  Links become w:hyperlink elements with an external relationship and the document's "Hyperlink" character style
  (created, with Word's look, if the document does not have it).
- Table cells hold either plain text (a string) or rich text ({"html": "..."}, from the cell editor): set_cell_value
  writes either kind into a cell, cell_value reads a cell back, as rich text only if it contains links.
"""
from copy import deepcopy
from html import escape

from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml.ns import qn
from docx.oxml.shared import OxmlElement
from docx.shared import RGBColor
from docx.table import _Cell
from docx.text.run import Run

from ..compilers.docx_utils import set_cell_text
from ..i18n import t
from .html_converter import HtmlBlock, HtmlRun, html_to_blocks

HYPERLINK_STYLE = "Hyperlink"
# Color of the "Hyperlink" style created in documents without it (Word's default).
LINK_COLOR = RGBColor(0x05, 0x63, 0xC1)


def parse_cell_value(value):
    """Validates a cell of the request: a string or {"html": string}; raises ValueError otherwise."""
    if isinstance(value, str):
        return value
    if isinstance(value, dict) and isinstance(value.get("html"), str) and set(value) == {"html"}:
        return {"html": value["html"]}
    raise ValueError(t("errors.cellValue"))


def cell_plain_text(value) -> str:
    """Text of a cell value, without formatting (e.g. for sorting or the minimum checks)."""
    if isinstance(value, str):
        return value
    return "\n".join("".join(run.text for run in block.runs) for block in html_to_blocks(value["html"]))


def html_run_elements(paragraph, html_runs: list[HtmlRun], base_properties=None) -> list:
    """
    Elements (w:r, w:hyperlink) of the runs, not yet placed in the paragraph. base_properties (w:rPr) is copied in
    every run, e.g. the formatting of the keyword they replace; bold and the like of the editor are added to it.
    """
    elements = []
    for html_run in html_runs:
        if html_run.text == "\n":
            # A line break goes at the end of the previous run (or in a run of its own at the beginning).
            last = elements[-1] if elements else None
            if last is not None and last.tag == qn("w:hyperlink"):
                last = last[-1]
            if last is None:
                last = _new_run(paragraph, "", base_properties)._r
                elements.append(last)
            last.append(OxmlElement("w:br"))
            continue
        run = _new_run(paragraph, html_run.text, base_properties)
        if html_run.bold:
            run.bold = True
        if html_run.italic:
            run.italic = True
        if html_run.underline:
            run.underline = True
        if html_run.strike:
            run.font.strike = True
        if not html_run.link:
            elements.append(run._r)
            continue
        _style_link(paragraph, run)
        previous = elements[-1] if elements else None
        if previous is not None and previous.tag == qn("w:hyperlink") and previous.get("_address") == html_run.link:
            previous.append(run._r)
        else:
            hyperlink = OxmlElement("w:hyperlink")
            relationship_id = paragraph.part.relate_to(html_run.link, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
            hyperlink.set(qn("r:id"), relationship_id)
            hyperlink.set(qn("w:history"), "1")
            # Temporary mark to join the next runs of the same link; removed below.
            hyperlink.set("_address", html_run.link)
            hyperlink.append(run._r)
            elements.append(hyperlink)
    for element in elements:
        if element.tag == qn("w:hyperlink"):
            element.attrib.pop("_address", None)
    return elements


def append_html_runs(paragraph, html_runs: list[HtmlRun], base_properties=None) -> None:
    """Adds the runs at the end of the paragraph (see html_run_elements)."""
    for element in html_run_elements(paragraph, html_runs, base_properties):
        paragraph._p.append(element)


def set_cell_value(cell: _Cell, value) -> None:
    """
    Writes a cell value: plain text keeps the formatting of the first paragraph and run of the cell (set_cell_text);
    rich text writes a paragraph per block of the HTML, with the properties of the first paragraph and the
    formatting of its first run.
    """
    if isinstance(value, str):
        set_cell_text(cell, value)
        return
    first, *others = cell.paragraphs
    for paragraph in others:
        paragraph._p.getparent().remove(paragraph._p)
    first_run = first._p.find(qn("w:r"))
    base_properties = first_run.find(qn("w:rPr")) if first_run is not None else None
    base_properties = deepcopy(base_properties) if base_properties is not None else None
    for child in list(first._p):
        if child.tag != qn("w:pPr"):
            first._p.remove(child)
    blocks = html_to_blocks(value["html"]) or [HtmlBlock(runs=[HtmlRun("")])]
    paragraph = first
    for index, block in enumerate(blocks):
        if index > 0:
            new_paragraph = deepcopy(first._p)
            for child in list(new_paragraph):
                if child.tag != qn("w:pPr"):
                    new_paragraph.remove(child)
            paragraph._p.addnext(new_paragraph)
            paragraph = cell.paragraphs[index]
        runs = list(block.runs)
        if block.list_type:
            # Cells have no list numbering: the item gets a text prefix.
            runs.insert(0, HtmlRun("• " if block.list_type == "bullet" else f"{block.list_index}. "))
        append_html_runs(paragraph, runs, base_properties)


def cell_value(cell: _Cell):
    """Value of an existing cell: rich text ({"html"}) if it contains web links, so editing keeps them; else text."""
    from .paragraph_reader import has_web_links, paragraph_html

    if not any(has_web_links(paragraph) for paragraph in cell.paragraphs):
        return cell.text
    return {"html": "".join(paragraph_html(paragraph) for paragraph in cell.paragraphs)}


def html_inline_runs(html: str) -> list[HtmlRun]:
    """
    Runs of the HTML as a single piece of a paragraph (e.g. in place of a keyword): its paragraphs are separated by
    line breaks and list items get a text prefix.
    """
    runs = []
    for index, block in enumerate(html_to_blocks(html)):
        if index > 0:
            runs.append(HtmlRun("\n"))
        if block.list_type:
            runs.append(HtmlRun("• " if block.list_type == "bullet" else f"{block.list_index}. "))
        runs.extend(block.runs)
    return runs


def html_escape_text(text: str) -> str:
    """Plain text as HTML of a paragraph of the editor (e.g. a plain cell opened in the cell editor)."""
    return "".join(f"<p>{escape(line)}</p>" for line in text.split("\n"))


def _new_run(paragraph, text: str, base_properties) -> Run:
    """New run with the text, not placed in the paragraph (python-docx needs the paragraph as parent)."""
    element = OxmlElement("w:r")
    if base_properties is not None:
        element.append(deepcopy(base_properties))
    run = Run(element, paragraph)
    if text:
        run.text = text
    return run


def _style_link(paragraph, run: Run) -> None:
    """
    Gives the run the "Hyperlink" character style of the document; in documents without it (e.g. created from
    python-docx's default template) the style is added, blue and underlined as in Word.
    """
    styles = paragraph.part.package.main_document_part.styles
    style = next(
        (style for style in styles if style.type == WD_STYLE_TYPE.CHARACTER
         and (style.style_id == HYPERLINK_STYLE or style.name == HYPERLINK_STYLE)),
        None,
    )
    if style is None:
        style = styles.add_style(HYPERLINK_STYLE, WD_STYLE_TYPE.CHARACTER)
        style.font.color.rgb = LINK_COLOR
        style.font.underline = True
        style.unhide_when_used = True
    run.style = style
