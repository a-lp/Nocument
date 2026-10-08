import base64
import binascii
from io import BytesIO
import os
import tempfile

from docx import Document
from docx.shared import Cm
from flask import Flask, jsonify, request
from werkzeug.utils import secure_filename

from .analyzers.analyzer_interface import IAnalyzer
from .analyzers.keyword_analyzer import DEFAULT_KEYWORD_PATTERN, KeywordAnalyzer
from .analyzers.section_analyzer import SectionAnalyzer
from .analyzers.style_analyzer import StyleAnalyzer
from .builder.contents import NewHeading, content_from_dict
from .builder.rich_text import parse_cell_value
from .builder.document_builder import ContentSequence, DocumentBuilder, block_insertion_point
from .builder.styles import adapt_missing_styles
from .builder.docx_files import template_to_document
from .compilers.compiler_interface import ICompiler
from .content_blocks.api import create_content_blocks_blueprint
from .content_blocks.repository import IContentBlockRepository, InMemoryContentBlockRepository
from .content_blocks.sqlite_repository import SqliteContentBlockRepository
from .compilers.document_compiler import DocumentCompiler
from .compilers.image_compiler import ImageCompiler
from .compilers.keyword_compiler import KeywordCompiler
from .compilers.table_compiler import TableCompiler
from .i18n import LANGUAGES, set_language, t
from .logging_setup import configure_logging
from .plugins.api import create_plugins_blueprint
from .plugins.registry import PluginRegistry
from .renderers.registry import RENDERERS, get_renderer, renderer_availability
from .renderers.render_interface import IRender
from .settings import environment_variables, load_settings, save_settings
from .version import backend_version
from .templates.api import create_templates_blueprint, docx_name, find_template
from .templates.repository import ITemplateRepository, InMemoryTemplateRepository
from .templates.repository import RepositoryError as TemplateRepositoryError
from .templates.sqlite_repository import SqliteTemplateRepository

# Before creating the app: Flask then uses the shared handlers instead of adding its own. With python -m src.app the
# reloader starts this module twice: the watcher process (WERKZEUG_RUN_MAIN not set yet) only logs to the console.
configure_logging(log_to_file=not (__name__ == "__main__" and os.getenv("WERKZEUG_RUN_MAIN") != "true"))
app = Flask(__name__)
# The body of /api/compile carries the document and the images in base64.
app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024

# ContentBlock and template repositories: MongoDB if MONGODB_URI is set (Docker Compose); otherwise a SQLite file,
# which needs no database server (NOCUMENT_DATABASE, default instance/nocument.sqlite3), or memory with
# NOCUMENT_STORAGE=memory (the data is lost when the backend restarts).
content_block_repository: IContentBlockRepository
template_repository: ITemplateRepository
if os.getenv("MONGODB_URI"):
    # Local import: pymongo is only needed with MongoDB.
    from .content_blocks.mongo_repository import MongoContentBlockRepository
    from .templates.mongo_repository import MongoTemplateRepository

    content_block_repository = MongoContentBlockRepository.from_uri(
        os.environ["MONGODB_URI"], os.getenv("MONGODB_DATABASE", "nocument")
    )
    template_repository = MongoTemplateRepository.from_uri(
        os.environ["MONGODB_URI"], os.getenv("MONGODB_DATABASE", "nocument")
    )
elif os.getenv("NOCUMENT_STORAGE", "").strip().lower() == "memory":
    app.logger.warning("NOCUMENT_STORAGE=memory: ContentBlocks and templates are kept in memory and lost on restart.")
    content_block_repository = InMemoryContentBlockRepository()
    template_repository = InMemoryTemplateRepository()
else:
    database_path = os.getenv("NOCUMENT_DATABASE") or os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "instance", "nocument.sqlite3"
    )
    content_block_repository = SqliteContentBlockRepository(database_path)
    template_repository = SqliteTemplateRepository(database_path)
app.logger.info("ContentBlocks saved with %s", type(content_block_repository).__name__)
app.register_blueprint(create_content_blocks_blueprint(content_block_repository))
app.register_blueprint(create_templates_blueprint(template_repository))

# Plugins: scripts in the plugins folder of the project (or NOCUMENT_PLUGINS_DIR). Installing them from the web app
# runs the uploaded Python code on the server: NOCUMENT_PLUGIN_UPLOAD=false allows only copying them by hand.
PROJECT_DIRECTORY = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
plugin_registry = PluginRegistry(os.getenv("NOCUMENT_PLUGINS_DIR") or os.path.join(PROJECT_DIRECTORY, "plugins"))
app.register_blueprint(create_plugins_blueprint(
    plugin_registry, os.getenv("NOCUMENT_PLUGIN_UPLOAD", "true").strip().lower() not in ("false", "0", "no", "off")
))


@app.before_request
def select_language():
    """Messages of the request in the language of the web app, sent by the frontend in Accept-Language."""
    set_language(request.accept_languages.best_match(LANGUAGES))


def renderer() -> IRender:
    """PDF renderer chosen in the settings (or the first available one, see renderers.registry.get_renderer)."""
    return get_renderer(load_settings()["renderer"])


def decode_base64(value, field: str) -> bytes:
    """Decodes a base64 string; raises ValueError with a message naming field if it is not valid."""
    if not isinstance(value, str):
        raise ValueError(t("errors.base64NotString", field=field))
    try:
        return base64.b64decode(value, validate=True)
    except binascii.Error as error:
        raise ValueError(t("errors.base64Invalid", field=field)) from error


def _analyze_document(document_path: str) -> list[str]:
    """Runs the analysis pipeline on the document and returns the keywords found."""
    keyword_analyzer = KeywordAnalyzer(DEFAULT_KEYWORD_PATTERN)
    # For now the sections stay in section_analyzer.contents and are not sent to the frontend.
    section_analyzer = SectionAnalyzer()
    pipeline: list[IAnalyzer] = [keyword_analyzer, section_analyzer]

    document = Document(document_path)
    for analyzer in pipeline:
        analyzer.analyze(document)
    return keyword_analyzer.keywords


def build_compiler(class_name, value) -> ICompiler:
    """Creates the compiler named by the "class" field of the request, validating its value."""
    if class_name == "KeywordCompiler":
        # Plain text, or {"html": ...} for rich text with hyperlinks.
        if not isinstance(value, str) and not (isinstance(value, dict) and isinstance(value.get("html"), str)):
            raise ValueError(t("errors.keywordCompilerValue"))
        return KeywordCompiler(value if isinstance(value, str) else {"html": value["html"]})
    if class_name == "ImageCompiler":
        # The image in base64, or {"data": base64, "width": cm} to give it a width (same field as the Builder's image).
        data, width = (value.get("data"), value.get("width")) if isinstance(value, dict) else (value, None)
        if width is not None and (isinstance(width, bool) or not isinstance(width, (int, float)) or not 0 < width <= 100):
            raise ValueError(t("errors.imageWidth"))
        return ImageCompiler(BytesIO(decode_base64(data, t("fields.imageCompilerValue"))), width=Cm(width) if width else None)
    if class_name == "TableCompiler":
        # The value is the list of rows or {"rows": [...], "column_widths": [...]}.
        rows, column_widths = (value.get("rows"), value.get("column_widths")) if isinstance(value, dict) else (value, None)
        if not isinstance(rows, list) or not all(isinstance(row, list) for row in rows):
            raise ValueError(t("errors.tableCompilerRows"))
        # Cells: plain text, or {"html": ...} for rich text with hyperlinks.
        rows = [[parse_cell_value(cell) for cell in row] for row in rows]
        if column_widths is not None and (
            not isinstance(column_widths, list) or not column_widths or not all(
                isinstance(width, (int, float)) and not isinstance(width, bool) and width > 0 for width in column_widths
            )
        ):
            raise ValueError(t("errors.tableCompilerWidths"))
        return TableCompiler(rows, column_widths)
    raise ValueError(t("errors.unsupportedCompiler", name=class_name))


def parse_keywords(keywords) -> list[tuple[str, ICompiler]]:
    """Validates the keyword list of the /api/compile body and creates the requested compiler for each one."""
    if not isinstance(keywords, list):
        raise ValueError(t("errors.invalidKeywords"))
    parsed = []
    for item in keywords:
        if not isinstance(item, dict) or not isinstance(item.get("keyword"), str) or not item["keyword"]:
            raise ValueError(t("errors.invalidKeywords"))
        parsed.append((item["keyword"], build_compiler(item.get("class"), item.get("value"))))
    return parsed


@app.get("/api/hello")
def hello():
    """Test endpoint used by Home to check that the backend answers."""
    return jsonify(message="HUehuehue")


def _settings_response():
    """
    Current settings with the available options and the environment variables (read-only, see
    settings.environment_variables): {"renderer", "renderers": [{"id", "available"}], "languages", "environment"}.
    """
    available = renderer_availability()
    return jsonify(
        renderer=load_settings()["renderer"],
        renderers=[{"id": renderer_id, "available": available[renderer_id]} for renderer_id in RENDERERS],
        languages=LANGUAGES,
        environment=environment_variables(),
    )


@app.get("/api/version")
def get_version():
    """Version of the backend: {"commit", "date"} of its git commit (see version.py), null if unknown."""
    return jsonify(backend_version())


@app.get("/api/settings")
def get_settings():
    """Settings of the web app saved on the server (see settings.py). Response: _settings_response."""
    return _settings_response()


@app.put("/api/settings")
def update_settings():
    """
    Changes the server settings. Body JSON: {"renderer": "libreoffice" | "word"}; only renderers available on this
    machine can be chosen. Response: _settings_response.
    """
    payload = request.get_json(silent=True)
    if not isinstance(payload, dict):
        return jsonify(error=t("errors.invalidJson")), 400
    changes = {}
    if "renderer" in payload:
        renderer_id = payload["renderer"]
        if renderer_id not in RENDERERS:
            return jsonify(error=t("errors.unknownRenderer", renderer=renderer_id)), 400
        if not renderer_availability()[renderer_id]:
            return jsonify(error=t("errors.rendererUnavailable", renderer=t(f"settings.renderer.names.{renderer_id}"))), 400
        changes["renderer"] = renderer_id
    save_settings(changes)
    return _settings_response()


@app.post("/api/render-pdf")
def render_pdf():
    """
    Converts the uploaded document to PDF and extracts its keywords.
    Response JSON: {"pdf": "<base64>", "keywords": ["<REV>", ...]}
    """
    uploaded_file = request.files.get("file")
    if not uploaded_file or not uploaded_file.filename.lower().endswith(".docx"):
        return jsonify(error=t("errors.uploadDocx")), 400

    safe_filename = secure_filename(uploaded_file.filename) or "document.docx"

    with tempfile.TemporaryDirectory() as temp_directory:
        input_path = os.path.join(temp_directory, safe_filename)
        uploaded_file.save(input_path)
        try:
            keywords = _analyze_document(input_path)
        except Exception:
            # Without the analysis the document can still be viewed: the user can add keywords by hand.
            app.logger.exception("Document analysis failed")
            keywords = []
        try:
            output_path = renderer().render_pdf(os.path.abspath(input_path), temp_directory)
            with open(output_path, "rb") as pdf_file:
                pdf_content = pdf_file.read()
        except Exception:
            app.logger.exception("Word to PDF conversion failed")
            return jsonify(error=t("errors.pdfConversion")), 500

    return jsonify(pdf=base64.b64encode(pdf_content).decode("ascii"), keywords=keywords)


@app.post("/api/parse-csv")
def parse_csv():
    """
    Reads the uploaded CSV to import it into a table.
    Response JSON: {"rows": [["...", ...], ...], "has_header": true}
    """
    uploaded_file = request.files.get("file")
    if not uploaded_file or not uploaded_file.filename.lower().endswith(".csv"):
        return jsonify(error=t("errors.uploadCsv")), 400

    try:
        rows, has_header = TableCompiler.read_csv(uploaded_file.read())
    except Exception:
        app.logger.exception("CSV parsing failed")
        return jsonify(error=t("errors.csvUnreadable")), 400
    if not rows:
        return jsonify(error=t("errors.csvEmpty")), 400

    return jsonify(rows=rows, has_header=has_header)


@app.post("/api/compile")
def compile_document():
    """
    Body JSON:
    {
      "file": {"name": "document.docx", "content": "<base64>"},
      "keywords": [
        {"keyword": "<REV>", "class": "KeywordCompiler", "value": "B"},
        {"keyword": "<LOGO>", "class": "ImageCompiler", "value": "<base64>"},
        {"keyword": "<TRCR>", "class": "TableCompiler", "value": [["101", "Fix"], ["102", "Export"]]}
      ]
    }
    Response JSON: {"pdf": "<base64>", "docx": "<base64>"} (compiled PDF and Word document)
    """
    payload = request.get_json(silent=True)
    if not isinstance(payload, dict):
        return jsonify(error=t("errors.invalidJson")), 400

    file = payload.get("file")
    if not isinstance(file, dict) or not str(file.get("name", "")).lower().endswith(".docx"):
        return jsonify(error=t("errors.uploadDocx")), 400

    try:
        document_content = decode_base64(file.get("content"), t("fields.fileContent"))
        keywords = parse_keywords(payload.get("keywords", []))
    except ValueError as error:
        return jsonify(error=str(error)), 400

    safe_filename = secure_filename(file["name"]) or "document.docx"

    with tempfile.TemporaryDirectory() as temp_directory:
        input_path = os.path.join(temp_directory, safe_filename)
        with open(input_path, "wb") as input_file:
            input_file.write(document_content)

        document_compiler = DocumentCompiler(input_path)
        for keyword, compiler in keywords:
            document_compiler.add_keyword(keyword, compiler)

        try:
            if not document_compiler.compile():
                return jsonify(error=t("errors.cannotOpenDocument")), 400
        except ValueError as error:
            # Content errors of the compilers (e.g. TableCompiler on a keyword outside a table).
            return jsonify(error=str(error)), 400
        except Exception:
            app.logger.exception("Document compilation failed")
            return jsonify(error=t("errors.compileFailed")), 500

        try:
            output_path = renderer().render_pdf(os.path.abspath(input_path), temp_directory)
            with open(output_path, "rb") as pdf_file:
                pdf_content = pdf_file.read()
        except Exception:
            app.logger.exception("Compiled document to PDF conversion failed")
            return jsonify(error=t("errors.compiledPdfConversion")), 500

        with open(input_path, "rb") as docx_file:
            docx_content = docx_file.read()

    return jsonify(
        pdf=base64.b64encode(pdf_content).decode("ascii"),
        docx=base64.b64encode(docx_content).decode("ascii"),
    )


def _builder_response(docx_content: bytes, filename: str, warnings: list[str] | None = None):
    """
    Analyzes the Builder document after opening it or after a change. Nothing slow happens here: the tables of
    contents (updated by LibreOffice) and the PDF are produced only when the finished document is asked for
    (/api/builder/export, for Preview PDF, Save PDF and Save Word), so every change is applied at once.
    Response JSON: {"name", "docx": "<base64>", "structure": [...], "styles": {...}, "warnings": [...]}
    (warnings: notes about the operation, e.g. styles missing from the document).
    """
    section_analyzer = SectionAnalyzer()
    style_analyzer = StyleAnalyzer()
    document = Document(BytesIO(docx_content))
    for analyzer in (section_analyzer, style_analyzer):
        analyzer.analyze(document)

    return jsonify(
        name=filename,
        docx=base64.b64encode(docx_content).decode("ascii"),
        structure=[content.to_dict() for content in section_analyzer.contents],
        styles=style_analyzer.to_dict(),
        warnings=warnings or [],
    )


@app.post("/api/builder/open")
def builder_open():
    """Opens a document (.docx) or a template (.dotx, converted to a document) in the Builder. Response: _builder_response."""
    uploaded_file = request.files.get("file")
    filename = uploaded_file.filename.lower() if uploaded_file else ""
    if not filename.endswith((".docx", ".dotx")):
        return jsonify(error=t("errors.uploadDocxOrDotx")), 400

    content = uploaded_file.read()
    try:
        if filename.endswith(".dotx"):
            content = template_to_document(content)
        Document(BytesIO(content))
    except Exception:
        app.logger.exception("Builder document opening failed")
        return jsonify(error=t("errors.cannotOpenDocument")), 400

    base_name = os.path.splitext(secure_filename(uploaded_file.filename))[0] or "document"
    return _builder_response(content, f"{base_name}.docx")


@app.post("/api/builder/new")
def builder_new():
    """Creates an empty document (default Word template) in the Builder. Response: _builder_response."""
    output = BytesIO()
    Document().save(output)
    return _builder_response(output.getvalue(), docx_name(t("builder.untitledDocument")))


@app.post("/api/builder/template/<template_id>")
def builder_from_template(template_id: str):
    """Creates a document in the Builder from a saved template, named after it. Response: _builder_response."""
    try:
        template = find_template(template_repository, template_id)
    except TemplateRepositoryError:
        app.logger.exception("Template repository error")
        return jsonify(error=t("errors.templatesDatabaseUnavailable")), 503
    if template is None:
        return jsonify(error=t("errors.templateNotFound")), 404
    return _builder_response(template.docx, docx_name(template.name))


def _builder_document(payload) -> tuple[Document, str]:
    """Reads the document {"file": {"name", "content"}} from the JSON body of a Builder route; raises ValueError."""
    if not isinstance(payload, dict):
        raise ValueError(t("errors.invalidJson"))
    file = payload.get("file")
    if not isinstance(file, dict) or not str(file.get("name", "")).lower().endswith(".docx"):
        raise ValueError(t("errors.invalidDocument"))
    try:
        document = Document(BytesIO(decode_base64(file.get("content"), t("fields.fileContent"))))
    except ValueError:
        raise
    except Exception as error:
        raise ValueError(t("errors.cannotOpenDocument")) from error
    return document, secure_filename(file["name"]) or "document.docx"


def _builder_target(payload) -> tuple[int, int]:
    """Reads {"target": {"id", "anchor"}}: the blocks taken by the element to edit or delete."""
    return _parse_target(payload.get("target"))


def _parse_target(target) -> tuple[int, int]:
    """Blocks (id, anchor) of a structure element; raises ValueError if they are not valid."""
    values = (target.get("id"), target.get("anchor")) if isinstance(target, dict) else (None, None)
    if not all(isinstance(value, int) and not isinstance(value, bool) for value in values):
        raise ValueError(t("errors.invalidTarget"))
    return values


def _parse_layout(target) -> int | None:
    """Position "layout" of an element inside a layout table (see SectionAnalyzer), or None."""
    layout = target.get("layout") if isinstance(target, dict) else None
    if layout is not None and (isinstance(layout, bool) or not isinstance(layout, int)):
        raise ValueError(t("errors.invalidTarget"))
    return layout


def _run_builder_edit(edit):
    """
    Runs edit(payload, builder) on the document of the request and returns the updated document
    (_builder_response) with the warnings returned by edit, if any; validation errors become 400.
    """
    payload = request.get_json(silent=True)
    try:
        document, filename = _builder_document(payload)
        warnings = edit(payload, DocumentBuilder(document))
    except ValueError as error:
        return jsonify(error=str(error)), 400
    except Exception:
        app.logger.exception("Builder edit failed")
        return jsonify(error=t("errors.updateFailed")), 500

    output = BytesIO()
    document.save(output)
    return _builder_response(output.getvalue(), filename, warnings)


@app.post("/api/builder/insert")
def builder_insert():
    """
    Inserts a content in the Builder document.
    Body JSON:
    {
      "file": {"name": "document.docx", "content": "<base64>"},
      "after": 12,            // id of the block to insert after (anchor in the structure), null = beginning
      "content": {"type": "heading", "text": "Scope", "level": 1, "style": "Heading1", "formatting": {...}}
    }
    Instead of "content" there can be "contents": [...], several contents inserted together and in order (e.g. the
    elements of a saved ContentBlock): their styles missing from the document are replaced with the default ones,
    with a warning in "warnings".
    Contents with headings are inserted as a block (see block_insertion_point): if after is inside a section, they
    go after the contents and subsections that follow it, so their headings do not take them over.
    Content types: heading (text, level), paragraph (html), image (data), table (rows, caption,
    caption_position, header, alignment); all with an optional style, heading and paragraph with optional formatting.
    Response: _builder_response of the updated document.
    """
    def insert(payload, builder):
        after = payload.get("after")
        if after is not None and (isinstance(after, bool) or not isinstance(after, int)):
            raise ValueError(t("errors.invalidInsertPosition"))
        if "contents" not in payload:
            content = content_from_dict(payload.get("content"))
            builder.insert(_block_position(builder, after, [content]), content)
            return []
        items = payload["contents"]
        if not isinstance(items, list) or not items:
            raise ValueError(t("errors.insertContentsNotList"))
        contents = [content_from_dict(item) for item in items]
        missing = adapt_missing_styles(builder.document, contents)
        builder.insert(_block_position(builder, after, contents), ContentSequence(contents))
        return [t("warnings.missingStyles", styles=", ".join(missing))] if missing else []

    return _run_builder_edit(insert)


def _block_position(builder: DocumentBuilder, after: int | None, contents: list) -> int | None:
    """Insertion point of the contents: after, moved past what their headings would take over (if they have any)."""
    levels = [content.level for content in contents if isinstance(content, NewHeading)]
    if not levels or (after is not None and not 0 <= after < len(builder.blocks)):
        # Without headings nothing is taken over; an invalid position is reported by DocumentBuilder.insert.
        return after
    return block_insertion_point(builder.document, after, min(levels))


@app.post("/api/builder/update")
def builder_update():
    """
    Replaces an existing element. Body JSON as /api/builder/insert, with "target": {"id", "anchor"}
    (from the structure) instead of "after"; for an element of a layout table also "layout".
    Response: _builder_response of the updated document.
    """
    def update(payload, builder):
        first, last = _builder_target(payload)
        layout = _parse_layout(payload.get("target"))
        content = content_from_dict(payload.get("content"))
        if layout is None:
            builder.replace(first, last, content)
        else:
            builder.replace_layout(first, layout, content)

    return _run_builder_edit(update)


@app.post("/api/builder/move")
def builder_move():
    """
    Moves an existing element (a heading together with its section) after another one.
    Body JSON: {"file": {...}, "target": {"id", "anchor"}, "after": 12}, where target are the blocks to move
    (from id to anchor included) and after is the anchor of the element to move them after (null = beginning).
    A moved heading goes as a block (see block_insertion_point): after the contents and subsections that follow after,
    so it does not take them over.
    Response: _builder_response of the updated document.
    """
    def move(payload, builder):
        after = payload.get("after")
        if after is not None and (isinstance(after, bool) or not isinstance(after, int)):
            raise ValueError(t("errors.invalidMovePosition"))
        first, last = _builder_target(payload)
        level = builder.heading_level_at(first)
        if level is not None and (after is None or 0 <= after < len(builder.blocks)) and not (
                after is not None and first <= after <= last):
            after = block_insertion_point(builder.document, after, level, (first, last))
        builder.move(first, last, after)

    return _run_builder_edit(move)


@app.post("/api/builder/delete")
def builder_delete():
    """
    Deletes one or more elements. Body JSON: {"file": {...}, "target": {"id", "anchor"}} or, for several elements
    together, {"file": {...}, "targets": [{"id", "anchor"}, ...]}. An element of a layout table also gives
    "layout" and deletes only itself. Response: _builder_response.
    """
    def delete(payload, builder):
        targets = payload["targets"] if "targets" in payload else [payload.get("target")]
        if not isinstance(targets, list) or not targets:
            raise ValueError(t("errors.deleteTargetsNotList"))
        ranges, layout_items = [], []
        for target in targets:
            first, last = _parse_target(target)
            layout = _parse_layout(target)
            if layout is None:
                ranges.append((first, last))
            else:
                layout_items.append((first, layout))
        builder.delete_many(ranges, layout_items)

    return _run_builder_edit(delete)


@app.post("/api/builder/export")
def builder_export():
    """
    Finished Builder document, for Preview PDF, Save PDF and Save Word (and before it leaves the Builder): its tables
    of contents are updated (new headings, page numbers) and, with "pdf": true, it is converted to PDF.
    Body JSON: {"file": {"name": "document.docx", "content": "<base64>"}, "pdf": false}.
    Response JSON: {"docx": "<base64>"} with the updated document and, with pdf, "pdf": "<base64>".
    """
    payload = request.get_json(silent=True)
    try:
        _, filename = _builder_document(payload)
    except ValueError as error:
        return jsonify(error=str(error)), 400
    with_pdf = bool(payload.get("pdf"))

    with tempfile.TemporaryDirectory() as temp_directory:
        input_path = os.path.join(temp_directory, filename)
        with open(input_path, "wb") as input_file:
            input_file.write(decode_base64(payload["file"]["content"], t("fields.fileContent")))
        # A failed update leaves the old entries: the document is still usable.
        try:
            renderer().update_toc(input_path, temp_directory)
        except Exception:
            app.logger.exception("Table of contents update failed")
        with open(input_path, "rb") as input_file:
            docx_content = input_file.read()
        response = {"docx": base64.b64encode(docx_content).decode("ascii")}
        if with_pdf:
            try:
                output_path = renderer().render_pdf(os.path.abspath(input_path), temp_directory)
                with open(output_path, "rb") as pdf_file:
                    response["pdf"] = base64.b64encode(pdf_file.read()).decode("ascii")
            except Exception:
                app.logger.exception("Builder PDF conversion failed")
                return jsonify(error=t("errors.builderPdfConversion")), 500

    return jsonify(response)


if __name__ == "__main__":
    # Development server; the Vite proxy forwards /api to the same NOCUMENT_BACKEND_PORT (see web/vite.config.js).
    app.run(
        debug=True,
        port=int(os.getenv("NOCUMENT_BACKEND_PORT") or 5000),
        host=os.getenv("NOCUMENT_BACKEND_HOST") or "127.0.0.1",
    )
