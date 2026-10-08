from docx.document import Document
from docx.enum.style import WD_STYLE_TYPE

from ..builder.styles import is_caption_style
from ..compilers.docx_utils import style_heading_level
from .analyzer_interface import IAnalyzer


class StyleAnalyzer(IAnalyzer):
    """
    Extracts the visible styles of the document, split into heading, paragraph and table styles.
    Each style is {"id", "name"} (headings also have "level"); the template's default styles are given in
    defaults ("heading" is the id of the style of each level, usually "Heading N").
    """

    def __init__(self):
        """Initializes the style lists, filled by analyze."""
        self.headings: list[dict] = []
        self.paragraphs: list[dict] = []
        self.tables: list[dict] = []
        self.caption: str | None = None
        self.defaults: dict = {}

    def analyze(self, document: Document) -> None:
        """Reads the document styles, excluding hidden ones."""
        styles = [style for style in document.styles if not style.hidden]
        default_paragraph = document.styles.default(WD_STYLE_TYPE.PARAGRAPH)
        default_table = document.styles.default(WD_STYLE_TYPE.TABLE)

        self.headings, self.paragraphs, self.tables = [], [], []
        self.caption = None
        for style in styles:
            if style.type == WD_STYLE_TYPE.PARAGRAPH:
                level = style_heading_level(style)
                if level is not None:
                    self.headings.append({"id": style.style_id, "name": style.name, "level": level})
                else:
                    self.paragraphs.append({"id": style.style_id, "name": style.name, "quick": style.quick_style})
                if is_caption_style(style):
                    self.caption = style.style_id
            elif style.type == WD_STYLE_TYPE.TABLE:
                self.tables.append({"id": style.style_id, "name": style.name})

        self.headings.sort(key=lambda style: (style["level"], style["id"] != f"Heading{style['level']}", style["name"]))
        # First the default style, then those of Word's quick style gallery, then the others.
        default_paragraph_id = default_paragraph.style_id if default_paragraph else None
        self.paragraphs.sort(key=lambda style: (style["id"] != default_paragraph_id, not style["quick"], style["name"].lower()))
        self.tables.sort(key=lambda style: style["name"].lower())

        heading_defaults = {}
        for style in self.headings:
            heading_defaults.setdefault(style["level"], style["id"])
        # For new tables the grid is more useful than "Normal Table" (no borders), if the template defines it.
        table_ids = {style["id"] for style in self.tables}
        self.defaults = {
            "heading": heading_defaults,
            "paragraph": default_paragraph_id,
            "table": "TableGrid" if "TableGrid" in table_ids else (default_table.style_id if default_table else None),
        }

    def to_dict(self) -> dict:
        """JSON representation of the styles, sent to the frontend."""
        return {
            "headings": self.headings,
            "paragraphs": [{"id": style["id"], "name": style["name"]} for style in self.paragraphs],
            "tables": self.tables,
            "caption": self.caption,
            "defaults": self.defaults,
        }
